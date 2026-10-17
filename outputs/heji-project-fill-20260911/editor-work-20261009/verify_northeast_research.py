import json
import sys
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')
work = Path(__file__).parent
with (work / 'original.xlsx.before-northeast-research').open('rb') as file:
    before = load_workbook(file)
after = load_workbook(work / 'northeast-research.xlsx')
edits = json.loads((work / 'northeast_research_updates.json').read_text(encoding='utf-8'))
expected = {'地图项目样例':dict(edits['main']), '待核验':dict(edits['pending'])}
for address, addition in edits['append'].items():
    expected['地图项目样例'][address] = before['地图项目样例'][address].value + '\n' + addition
issues = []
assert before.sheetnames == after.sheetnames
style_props = ['font','fill','border','alignment','number_format','protection']
for old, new in zip(before, after):
    for row in old:
        for cell in row:
            if old.title == '待核验' and 15 <= cell.row <= 17 and cell.column <= 4:
                continue
            actual = new[cell.coordinate]
            if actual.value != expected[old.title].get(cell.coordinate, cell.value):
                issues.append(f'{old.title}!{cell.coordinate}: value')
            for prop in style_props:
                if copy(getattr(cell, prop)) != copy(getattr(actual, prop)):
                    issues.append(f'{old.title}!{cell.coordinate}: {prop}')
    for prop in ['freeze_panes','auto_filter','data_validations']:
        if getattr(old,prop) != getattr(new,prop):
            issues.append(f'{old.title}: {prop}')
    if str(old.merged_cells) != str(new.merged_cells):
        issues.append(f'{old.title}: merges')
    for col, dim in old.column_dimensions.items():
        for prop in ['width','hidden','min','max']:
            if getattr(dim,prop) != getattr(new.column_dimensions[col],prop):
                issues.append(f'{old.title}: {col} {prop}')
    assert list(old.tables) == list(new.tables)
    for name in old.tables:
        assert old.tables[name].ref == new.tables[name].ref
        assert old.tables[name].tableStyleInfo == new.tables[name].tableStyleInfo
    for row, dim in old.row_dimensions.items():
        height = 330 if old.title == '地图项目样例' and row == 10 else dim.height
        if old.title == '待核验' and 15 <= row <= 17:
            height = 150
        if height != new.row_dimensions[row].height:
            issues.append(f'{old.title}: {row} height')
for row, intended in enumerate(edits['new_pending'], 15):
    assert [after['待核验'].cell(row,col).value for col in range(1,5)] == intended
    for col in range(1,5):
        for prop in style_props:
            assert copy(getattr(before['待核验'].cell(12,col),prop)) == copy(getattr(after['待核验'].cell(row,col),prop))
assert len(set(after['地图项目样例'].cell(row,1).value for row in range(5,18))) == 13
assert after['地图项目样例']['C10'].value == before['地图项目样例']['C10'].value
assert after['地图项目样例']['D10'].value == before['地图项目样例']['D10'].value
report = {'projects':13,'changed_existing_cells':sum(map(len,expected.values())),'new_pending_rows':3,'issues':issues}
(work/'northeast-research-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
if issues:
    sys.exit(1)
