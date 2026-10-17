import json
import sys
from copy import copy
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')
work = Path(__file__).parent
with (work / 'original.xlsx.before-channels').open('rb') as source:
    before = load_workbook(source)
after = load_workbook(work / 'channels.xlsx')
edits = json.loads((work / 'channels_updates.json').read_text(encoding='utf-8'))
expected = {'地图项目样例': dict(edits['updates']), '待核验': dict(edits['pending_updates'])}
for address, addition in edits['sources'].items():
    expected['地图项目样例'][address] = before['地图项目样例'][address].value + '\n' + addition
for address in ['F10', 'H10', 'I10', 'K10']:
    expected['地图项目样例'][address] = '沙棘项目：' + before['地图项目样例'][address].value
issues = []
if before.sheetnames != after.sheetnames:
    issues.append('sheet names/order changed')
for old, new in zip(before, after):
    for row in old:
        for cell in row:
            current = new[cell.coordinate]
            if old.title == '待核验' and 12 <= cell.row < 12 + len(edits['pending_new']) and cell.column <= 4:
                continue
            intended = expected[old.title].get(cell.coordinate, cell.value)
            if current.value != intended:
                issues.append(f'{old.title}!{cell.coordinate}: unexpected value')
            for prop in ['font', 'fill', 'border', 'alignment', 'protection', 'number_format']:
                if copy(getattr(cell, prop)) != copy(getattr(current, prop)):
                    issues.append(f'{old.title}!{cell.coordinate}: {prop} changed')
    for prop in ['freeze_panes']:
        if getattr(old, prop) != getattr(new, prop):
            issues.append(f'{old.title}: {prop} changed')
    if str(old.merged_cells) != str(new.merged_cells):
        issues.append(f'{old.title}: merges changed')
    if old.sheet_view.showGridLines != new.sheet_view.showGridLines:
        issues.append(f'{old.title}: grid lines changed')
    for col, dim in old.column_dimensions.items():
        for prop in ['width', 'hidden', 'min', 'max']:
            if getattr(dim, prop) != getattr(new.column_dimensions[col], prop):
                issues.append(f'{old.title}: column {col} {prop} changed')
    for row, dim in old.row_dimensions.items():
        height = edits['row_heights'].get(str(row), dim.height) if old.title == '地图项目样例' else dim.height
        if old.title == '待核验' and 12 <= row < 12 + len(edits['pending_new']):
            height = 96
        if new.row_dimensions[row].height != height:
            issues.append(f'{old.title}: row {row} height changed unexpectedly')
    if list(old.tables) != list(new.tables):
        issues.append(f'{old.title}: table names changed')
    for name, table in old.tables.items():
        original = old.tables[name]
        current = new.tables[name]
        if original.ref != current.ref or original.tableStyleInfo != current.tableStyleInfo:
            issues.append(f'{old.title}: table {name} range/style changed')
for index, values in enumerate(edits['pending_new'], 12):
    actual = [after['待核验'].cell(index, col).value for col in range(1, 5)]
    if actual != values:
        issues.append(f'pending row {index} mismatch')
    for col in range(1, 5):
        original = before['待核验'].cell(11, col)
        current = after['待核验'].cell(index, col)
        for prop in ['font', 'fill', 'border', 'alignment', 'protection', 'number_format']:
            if copy(getattr(original, prop)) != copy(getattr(current, prop)):
                issues.append(f'new pending {index}/{col}: {prop} differs')
main = after['地图项目样例']
projects = [main.cell(row, 1).value for row in range(5, 18)]
assert len(set(projects)) == 13
assert main['C10'].value == before['地图项目样例']['C10'].value
assert main['D10'].value == before['地图项目样例']['D10'].value
report = {'projects': len(projects), 'channels': 7, 'changed_existing_cells': sum(map(len, expected.values())), 'new_pending_rows': len(edits['pending_new']), 'issues': issues}
(work / 'channels-verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False))
if issues:
    sys.exit(1)
