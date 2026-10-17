from pathlib import Path
from copy import copy
import io
import json
import sys
import zipfile
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
work = Path(__file__).resolve().parent
original = openpyxl.load_workbook(io.BytesIO((work / "original.xlsx.before-edit").read_bytes()))
edited = openpyxl.load_workbook(work / "edited.xlsx")
allowed = json.loads((work / "allowed_changes.json").read_text())
issues = []
changed = []
style_changes = []

if original.sheetnames != edited.sheetnames:
    issues.append("Sheet names/order changed")
for old in original:
    new = edited[old.title]
    for row in old:
        for cell in row:
            other = new[cell.coordinate]
            permitted = old.title == "地图项目样例" and cell.coordinate in allowed["cells"]
            if cell.value != other.value:
                changed.append(f"{old.title}!{cell.coordinate}")
                if not permitted:
                    issues.append(f"Unexpected value change {old.title}!{cell.coordinate}")
                elif cell.coordinate not in ("A1", "A2") and not other.value.startswith(cell.value + "\n"):
                    issues.append(f"Original text replaced {cell.coordinate}")
            if cell.value is not None:
                components = []
                for attribute in ("font", "fill", "border", "alignment", "protection", "number_format"):
                    if copy(getattr(cell, attribute)) != copy(getattr(other, attribute)):
                        components.append(attribute)
                if components:
                    style_changes.append({"cell": f"{old.title}!{cell.coordinate}", "components": components})
    if set(map(str, old.merged_cells.ranges)) != set(map(str, new.merged_cells.ranges)):
        issues.append(f"Merged ranges changed {old.title}")
    if old.freeze_panes != new.freeze_panes:
        issues.append(f"Freeze panes changed {old.title}")
    if old.sheet_view.showGridLines != new.sheet_view.showGridLines:
        issues.append(f"Gridlines changed {old.title}")
    if set(old.tables) != set(new.tables):
        issues.append(f"Tables changed {old.title}")
    for key, dim in old.column_dimensions.items():
        other = new.column_dimensions[key]
        if abs(dim.width - other.width) > 0.05 or dim.hidden != other.hidden:
            issues.append(f"Column dimensions changed {old.title}!{key}")
    if len(old.data_validations.dataValidation) != len(new.data_validations.dataValidation):
        issues.append(f"Data validations changed {old.title}")
    if len(old.conditional_formatting) != len(new.conditional_formatting):
        issues.append(f"Conditional formatting changed {old.title}")
    for number, dim in old.row_dimensions.items():
        permitted = old.title == "地图项目样例" and number in allowed["changedRowHeights"]
        if not permitted and dim.height != new.row_dimensions[number].height:
            issues.append(f"Unexpected row height change {old.title}!{number}")

main = edited["地图项目样例"]
names = [main.cell(row, 1).value for row in range(5, 21)]
if len(names) != 16 or len(set(names)) != 16 or any(name is None for name in names):
    issues.append("Expected 16 distinct projects/events")
if main.tables["MapProjects"].ref != "A4:L20":
    issues.append(f"Table range is {main.tables['MapProjects'].ref}, expected A4:L20")
for row in range(15, 21):
    if main.cell(row, 3).value is not None or main.cell(row, 4).value is not None:
        issues.append(f"Invented coordinates row {row}")
    if "资料提取：2026-10-09" not in main.cell(row, 12).value:
        issues.append(f"Missing extraction date row {row}")

result = {
    "issues": issues,
    "changed_existing_cells": changed,
    "original_style_changes": style_changes,
    "tables_before": [(t.name, t.ref) for t in original["地图项目样例"].tables.values()],
    "table_after": main.tables["MapProjects"].ref,
    "project_count": len(names),
    "pending_count": edited["待核验"].max_row - 1,
}
(work / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False))
if issues or style_changes:
    sys.exit(1)
