from copy import copy
from pathlib import Path
import io
import json
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
work = Path(__file__).resolve().parent
before = openpyxl.load_workbook(io.BytesIO((work / "original.xlsx.before-merge").read_bytes()))
after = openpyxl.load_workbook(work / "merged.xlsx")
issues = []
changes = []
main_changes = {1: {1}, 2: {1}, 5: {7, 9, 10, 11, 12}, 18: set(range(6, 13)), 19: {1, 9, 10, 11}}
pending_changes = {2: {2, 3, 4}, 8: {2, 3, 4}, 9: {1, 3}}
main_map = {**{r: r for r in range(1, 15)}, 18: 15, 19: 16, 20: 17}
pending_map = {**{r: r for r in range(1, 5)}, **{r: r - 3 for r in range(8, 15)}}

assert before.sheetnames == after.sheetnames
for name, mapping, allowed in [("地图项目样例", main_map, main_changes), ("待核验", pending_map, pending_changes)]:
    old, new = before[name], after[name]
    for old_row, new_row in mapping.items():
        for col in range(1, old.max_column + 1):
            a, b = old.cell(old_row, col), new.cell(new_row, col)
            if a.value != b.value:
                changes.append(f"{name}!{a.coordinate}->{b.coordinate}")
                if col not in allowed.get(old_row, set()):
                    issues.append(f"Unexpected content change: {name}!{a.coordinate}")
            if a.value is not None:
                for component in ("font", "fill", "border", "alignment", "protection", "number_format"):
                    if copy(getattr(a, component)) != copy(getattr(b, component)):
                        issues.append(f"Style change {name}!{a.coordinate}: {component}")
    for component in ("freeze_panes", "print_area", "print_title_rows"):
        if getattr(old, component) != getattr(new, component):
            issues.append(f"Changed {name} {component}")
    if set(map(str, old.merged_cells.ranges)) != set(map(str, new.merged_cells.ranges)):
        issues.append(f"Changed merges {name}")
    for key, dimension in old.column_dimensions.items():
        other = new.column_dimensions[key]
        if abs(dimension.width - other.width) > 0.05 or dimension.hidden != other.hidden:
            issues.append(f"Changed column {name}!{key}")
    if old.sheet_view.showGridLines != new.sheet_view.showGridLines:
        issues.append(f"Changed gridlines {name}")

main = after["地图项目样例"]
names = [main.cell(row, 1).value for row in range(5, 18)]
if len(names) != 13 or len(set(names)) != 13 or any(n is None for n in names):
    issues.append("Expected 13 distinct project/event rows")
if any(main.cell(row, col).value is not None for row in range(18, 21) for col in range(1, 13)):
    issues.append("Removed rows retained values")
if any(any(term in n for term in ("民勤人参果助农直播", "后陡门周边果农草莓采购", "杨梅")) for n in names):
    issues.append("Standalone merged/deleted projects remain")
if main["A16"].value != "西双版纳水果助销":
    issues.append("Fruit project name incorrect")
if "购买民勤当地农产品" not in main["J5"].value or "人参果" not in main["J5"].value:
    issues.append("Minqin purchase guidance missing")
if "草莓" not in main["I15"].value or "EV003" not in main["L15"].value:
    issues.append("Strawberry action/source not merged into A Nong Day")
table = main.tables["MapProjects"]
if table.ref != "A4:L17":
    issues.append("Incorrect table extent")
if copy(table.tableStyleInfo) != copy(before["地图项目样例"].tables["MapProjects"].tableStyleInfo):
    issues.append("Table style changed")
if [(c.id, c.name) for c in table.tableColumns] != [(c.id, c.name) for c in before["地图项目样例"].tables["MapProjects"].tableColumns]:
    issues.append("Table columns changed")
queue_names = [after["待核验"].cell(row, 1).value for row in range(2, 12)]
if any(n in ("民勤人参果直播", "周边果农草莓采购", "缙云／仙居杨梅采购", "西双版纳受灾水果") for n in queue_names):
    issues.append("Pending queue contains obsolete standalone names")
if "西双版纳水果助销" not in queue_names:
    issues.append("Pending queue fruit name not updated")
result = {"issues": issues, "project_count": len(names), "pending_count": len(queue_names), "table_range": table.ref, "expected_changed_cells": len(changes)}
(work / "merge-verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False))
if issues:
    sys.exit(1)
