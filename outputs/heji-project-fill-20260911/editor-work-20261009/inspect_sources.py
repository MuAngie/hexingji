from pathlib import Path
import json
import sys
import zipfile
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
root = Path(__file__).resolve().parents[3]
work = Path(__file__).resolve().parent
source = root / "reference/种地吧全四季_助农关系数据库_升级版.xlsx"
target = work.parent / "禾伙人助农公益地图_项目汇总_v3.xlsx"

source_book = openpyxl.load_workbook(source, read_only=True, data_only=True)
source_rows = {sheet.title: list(sheet.values) for sheet in source_book}
(work / "source_rows.json").write_text(json.dumps(source_rows, ensure_ascii=False, default=str), encoding="utf-8")

book = openpyxl.load_workbook(target, data_only=False)
baseline = {}
for sheet in book:
    baseline[sheet.title] = {
        "rows": sheet.max_row,
        "cols": sheet.max_column,
        "cells": {cell.coordinate: {"value": cell.value, "style_id": cell.style_id,
                  "format": cell.number_format, "hyperlink": cell.hyperlink.target if cell.hyperlink else None}
                  for row in sheet for cell in row if cell.value is not None},
        "merges": [str(r) for r in sheet.merged_cells.ranges],
        "freeze": sheet.freeze_panes,
        "filter": sheet.auto_filter.ref,
        "tables": list(sheet.tables),
        "dimensions": {key: {"width": dim.width, "hidden": dim.hidden} for key, dim in sheet.column_dimensions.items()},
        "row_heights": {str(key): dim.height for key, dim in sheet.row_dimensions.items()},
        "gridlines": sheet.sheet_view.showGridLines,
        "print_area": str(sheet.print_area),
        "print_titles": sheet.print_title_rows,
        "validation_count": len(sheet.data_validations.dataValidation),
        "conditional_formats": len(sheet.conditional_formatting),
    }
(work / "baseline.json").write_text(json.dumps(baseline, ensure_ascii=False, default=str), encoding="utf-8")
with zipfile.ZipFile(target) as archive:
    native_parts = [name for name in archive.namelist() if any(part in name for part in ("tables/", "drawings/", "pivot", "externalLink", "comments"))]
print(json.dumps({"target": str(target), "sheets": baseline, "native_parts": native_parts}, ensure_ascii=False, default=str))
